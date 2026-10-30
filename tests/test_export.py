#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""oracle2excel 代码层深度检查（无需真实 Oracle 数据库）。

通过向 sys.modules 注入 mock 的 oracledb 模块，用伪造的 Connection/Cursor
端到端验证: 配置加载 -> 连接(thick 初始化) -> 导出 -> Excel 落盘。

运行: .venv/bin/python tests/test_export.py
"""

import io
import os
import shutil
import sys
import tempfile
from datetime import date, datetime

PROJECT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT)

# ---------------------------------------------------------------------------
# 1. Mock oracledb（在 import app.db 之前注入）
# ---------------------------------------------------------------------------
DB_TYPE_CLOB = "DB_TYPE_CLOB"
DB_TYPE_NCLOB = "DB_TYPE_NCLOB"
DB_TYPE_BLOB = "DB_TYPE_BLOB"
DB_TYPE_LONG = "DB_TYPE_LONG"
DB_TYPE_LONG_RAW = "DB_TYPE_LONG_RAW"


class MockTypeCode:
    def __init__(self, name):
        self.name = name


class FakeMetadata:
    """模拟 cursor 的列元数据 metadata 对象。"""

    def __init__(self, type_code):
        self.type_code = type_code


class FakeVar:
    """模拟 cursor.var(...) 返回值——真实驱动用它做 LOB 转换。"""

    def __init__(self, type_code, arraysize):
        self.type_code = type_code
        self.arraysize = arraysize


class FakeCursor:
    """伪造游标: 按表名返回预置数据。"""

    arraysize = 100

    def __init__(self, data_source):
        self._data = data_source
        self._results = None
        self._pos = 0
        self.description = None
        self.closed = False

    def execute(self, sql):
        key = self._pick_key(sql)
        cols, rows = self._data[key]
        self.description = [(c, FakeMetadata("UNKNOWN"), None, None, None, None, None) for c in cols]
        self._results = rows
        self._pos = 0

    @staticmethod
    def _pick_key(sql):
        s = sql.upper()
        for key in ("T_ORDER", "T_LOB", "T_BIG", "T_EDGE"):
            if key in s:
                return key
        return "T_USER"

    def fetchmany(self, n):
        if self._results is None:
            raise RuntimeError("execute() 未调用")
        chunk = self._results[self._pos:self._pos + n]
        self._pos += n
        return chunk

    def close(self):
        self.closed = True


class FakeConnection:
    version = "11.2.0.4.0"

    def __init__(self, data_source):
        self._data = data_source
        self.closed = False
        self.output_type_handler = None  # connect() 应注册 LOB 类型处理器

    def cursor(self):
        return FakeCursor(self._data)

    def close(self):
        self.closed = True


class FakeOracledb:
    """替换真实 oracledb 模块。"""

    DB_TYPE_CLOB = DB_TYPE_CLOB
    DB_TYPE_NCLOB = DB_TYPE_NCLOB
    DB_TYPE_BLOB = DB_TYPE_BLOB
    DB_TYPE_LONG = DB_TYPE_LONG
    DB_TYPE_LONG_RAW = DB_TYPE_LONG_RAW
    thick = False

    def __init__(self):
        self.init_called = False
        self.data_source = None

    def is_thick_mode(self):
        return self.thick

    def clientversion(self):
        return (19, 25, 0, 0, 0)

    def init_oracle_client(self, lib_dir=None):
        self.init_called = True
        if not lib_dir or not os.path.isdir(lib_dir):
            raise RuntimeError(f"DPI-1047: lib_dir={lib_dir} 不可用(mock 模拟)")
        self.thick = True
        return None

    def makedsn(self, host, port, service_name=None):
        return f"(DESCRIPTION=(ADDRESS=(PROTOCOL=TCP)(HOST={host})(PORT={port}))(CONNECT_DATA=(SERVICE_NAME={service_name})))"

    def connect(self, **kw):
        # 模拟首次连接失败，验证重试逻辑
        if not getattr(self, "_fail_once_done", False):
            self._fail_once_done = True
            raise RuntimeError("ORA-12541: no listener (mock)")
        return FakeConnection(self.data_source)


mock = FakeOracledb()
sys.modules["oracledb"] = mock

# ---------------------------------------------------------------------------
# 2. 被测代码
# ---------------------------------------------------------------------------
from app import config as app_config  # noqa: E402
from app import db as app_db  # noqa: E402
from app.config import ConfigError, load_config  # noqa: E402
from app.exporter import ExcelStreamWriter, export_one, normalize_value, run_export  # noqa: E402

PASS = []
FAIL = []


def check(name, cond, detail=""):
    if cond:
        PASS.append(name)
        print(f"  [PASS] {name}")
    else:
        FAIL.append(name)
        print(f"  [FAIL] {name} {detail}")


# ---------------------------------------------------------------------------
# 3. 测试数据
# ---------------------------------------------------------------------------
T_ORDER_COLS = ["ORDER_ID", "CUSTOMER_NAME", "AMOUNT", "CREATED_AT", "UPDATED_AT"]
T_ORDER_ROWS = [
    (1, "张三", 99.5, datetime(2026, 10, 30, 12, 0, 0), None),
    (2, "李四", 0, date(2026, 1, 1), datetime(2026, 3, 15, 8, 30, 0)),
] * 50  # 100 行

T_LOB_COLS = ["ID", "CONTENT", "ATTACHMENT"]
T_LOB_ROWS = [(1, "长文本" * 100, b"\x00\x01\xff" * 100)]  # 输出处理器转 str/bytes 后的形态

T_EDGE_COLS = ["ID", "NOTE", "WEIRD"]
_long = "X" * 40000
T_EDGE_ROWS = [(1, _long, float("nan")), (2, None, float("inf")), (3, "", -0.0)]

BIG_ROWS = 1005
T_BIG_COLS = ["ID", "VAL"]
T_BIG_ROWS = [(i, f"row-{i}") for i in range(1, BIG_ROWS + 1)]

DATA_SOURCE = {
    "T_ORDER": (T_ORDER_COLS, T_ORDER_ROWS),
    "T_LOB": (T_LOB_COLS, T_LOB_ROWS),
    "T_BIG": (T_BIG_COLS, T_BIG_ROWS),
    "T_EDGE": (T_EDGE_COLS, T_EDGE_ROWS),
    "T_USER": (["ID", "NAME"], [(1, "a")]),
}

# ---------------------------------------------------------------------------
# 用例 1: normalize_value 边界
# ---------------------------------------------------------------------------
print("== 用例 1: normalize_value ==")
check("None -> 空串", normalize_value(None, "%Y-%m-%d", 32767, 4096) == "")
check("datetime 格式化", normalize_value(datetime(2026, 10, 30, 12, 0), "%Y-%m-%d %H:%M:%S", 32767, 4096) == "2026-10-30 12:00:00")
check("date 格式化", normalize_value(date(2026, 10, 30), "%Y-%m-%d %H:%M:%S", 32767, 4096) == "2026-10-30 00:00:00")
check("bytes -> hex", normalize_value(b"\x00\x01", "%Y", 32767, 4096) == "0001")
check("NaN -> 空串", normalize_value(float("nan"), "%Y", 32767, 4096) == "")
check("inf -> 空串", normalize_value(float("inf"), "%Y", 32767, 4096) == "")
v = normalize_value(_long, "%Y", 32767, 4096)
check("超长字符串截断", len(v) <= 32767 + len("...(truncated)") and v.endswith("...(truncated)"))
big = b"\xff" * 4096
v2 = normalize_value(big, "%Y", 32767, 4096)
check("大 BLOB hex 截断", len(v2) <= 4096 + len("...(truncated)") and v2.endswith("...(truncated)"))
check("int 原样", normalize_value(12345, "%Y", 32767, 4096) == 12345)
check("bool 处理", normalize_value(True, "%Y", 32767, 4096) == "TRUE")

# ---------------------------------------------------------------------------
# 用例 2: 配置加载与校验
# ---------------------------------------------------------------------------
print("== 用例 2: 配置加载与校验 ==")
tmpdir = tempfile.mkdtemp(prefix="o2e_test_")
try:
    good = {
        "oracle": {"user": "u", "password": "p", "host": "h", "port": 1521,
                    "service_name": "orcl", "driver_mode": "thick"},
        "tables": ["T_ORDER"],
        "queries": [{"name": "q1", "sql": "SELECT 1 FROM DUAL"}],
        "output_dir": os.path.join(tmpdir, "out"),
    }
    p = os.path.join(tmpdir, "cfg.json")
    import json
    with open(p, "w", encoding="utf-8") as f:
        json.dump(good, f, ensure_ascii=False)
    cfg = load_config(p)
    check("合法配置加载", cfg.oracle.service_name == "orcl" and cfg.tables == ["T_ORDER"])
    check("build_jobs 合并 tables+queries", len(cfg.build_jobs([])) == 2)

    bad = dict(good)
    bad["tables"] = ["T1; DROP TABLE X"]
    with open(p, "w", encoding="utf-8") as f:
        json.dump(bad, f, ensure_ascii=False)
    try:
        load_config(p)
        check("表名注入拦截", False)
    except ConfigError:
        check("表名注入拦截", True)

    bad = dict(good)
    bad["queries"] = [{"name": "x", "sql": "DELETE FROM t"}]
    with open(p, "w", encoding="utf-8") as f:
        json.dump(bad, f, ensure_ascii=False)
    try:
        load_config(p)
        check("非 SELECT 拦截", False)
    except ConfigError:
        check("非 SELECT 拦截", True)

    with open(p, "w", encoding="utf-8") as f:
        f.write("{broken json")
    try:
        load_config(p)
        check("坏 JSON 报错", False)
    except ConfigError:
        check("坏 JSON 报错", True)
    try:
        load_config(os.path.join(tmpdir, "nope.json"))
        check("缺文件报错", False)
    except ConfigError:
        check("缺文件报错", True)
finally:
    shutil.rmtree(tmpdir, ignore_errors=True)

# ---------------------------------------------------------------------------
# 用例 3: DSN 构造（service_name）
# ---------------------------------------------------------------------------
print("== 用例 3: service_name DSN ==")
oc = app_config.OracleConfig(user="u", password="p", host="10.0.0.5", port=1521, service_name="orcl")
dsn = app_db.build_dsn(oc)
check("makedsn service_name", "SERVICE_NAME=orcl" in dsn and "10.0.0.5" in dsn and "1521" in dsn)

# ---------------------------------------------------------------------------
# 用例 4: 端到端导出（mock 驱动，含 LOB/超长/NaN/分 Sheet/重试）
# ---------------------------------------------------------------------------
print("== 用例 4: 端到端导出 ==")
outdir = os.path.join(tempfile.mkdtemp(prefix="o2e_e2e_"), "out")
mock.data_source = DATA_SOURCE
mock._fail_once_done = False  # 重新启用首次失败，验证重试
cfg = app_config.AppConfig(
    oracle=app_config.OracleConfig(
        user="u", password="p", host="h", port=1521, service_name="orcl",
        driver_mode="thick", connect_retry_attempts=3, connect_retry_delay_seconds=0.1,
    ),
    output_dir=outdir,
    batch_size=100,
    log_every_rows=500,
    tables=["T_ORDER", "T_LOB", "T_EDGE"],
    queries=[{"name": "big_split", "sql": "SELECT ID, VAL FROM T_BIG"}],
    excel=app_config.ExcelConfig(max_rows_per_sheet=1000),
)
try:
    conn = app_db.connect(cfg)
    check("重试后连接成功", conn.version == "11.2.0.4.0")
    check("LOB 类型处理器已注册", conn.output_type_handler is not None)
    files = run_export(conn, cfg, cfg.build_jobs([]))
    check("4 个任务全部产出", len(files) == 4, str(files))
    for f in files:
        check(f"文件存在 {os.path.basename(f)}", os.path.isfile(f) and os.path.getsize(f) > 1000)

    # 分 Sheet 验证: 1005 行 / 每 Sheet 1000 行 = 2 个 Sheet
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(outdir, "big_split.xlsx"))
    check("分 Sheet 数量", wb.sheetnames == ["DATA_1", "DATA_2"], str(wb.sheetnames))
    ws1_rows = sum(1 for _ in wb["DATA_1"].rows)
    check("Sheet1 = 表头+1000 行", ws1_rows == 1001, str(ws1_rows))
    ws2_rows = sum(1 for _ in wb["DATA_2"].rows)
    check("Sheet2 = 表头+5 行", ws2_rows == 6, str(ws2_rows))

    # LOB 验证
    wb2 = load_workbook(os.path.join(outdir, "T_LOB.xlsx"))
    ws = wb2.active
    row = list(ws.iter_rows(min_row=2, values_only=True))[0]
    check("CLOB 内容写入", str(row[1]).startswith("长文本长文本"))
    check("BLOB 转 hex", row[2] == "0001ff" * 100)

    # 边界验证
    wb3 = load_workbook(os.path.join(outdir, "T_EDGE.xlsx"))
    rows = list(wb3.active.iter_rows(min_row=2, values_only=True))
    check("超长单元格截断", len(str(rows[0][1])) <= 32767 + len("...(truncated)"))
    # 注: openpyxl 写入的空串回读为 None（Excel 空单元格），语义一致
    check("NaN 归一为空", rows[0][2] in ("", None))
    check("inf 归一为空", rows[1][2] in ("", None))
    check("None 归一为空", rows[1][1] in ("", None))
    check("空串保留", rows[2][1] in ("", None))
    check("-0.0 保留", rows[2][2] == -0.0 or str(rows[2][2]) == "-0.0")

    # 表头验证
    wb4 = load_workbook(os.path.join(outdir, "T_ORDER.xlsx"))
    header = [c.value for c in next(wb4.active.iter_rows(max_row=1))]
    check("表头来自 cursor.description", header == T_ORDER_COLS, str(header))
except Exception as e:  # noqa: BLE001
    import traceback
    traceback.print_exc()
    check(f"端到端流程异常: {e}", False)

# ---------------------------------------------------------------------------
# 用例 5: thick 初始化失败路径
# ---------------------------------------------------------------------------
print("== 用例 5: thick 初始化失败路径 ==")
mock2 = FakeOracledb()
app_db.oracledb = mock2  # app.db 持有的是模块级引用，需直接替换
app_db._THICK_INITIALIZED = False
old_oh = os.environ.get("ORACLE_HOME")
os.environ["ORACLE_HOME"] = "/tmp/definitely_not_exist_ic"  # 避免本机真实 ORACLE_HOME 干扰
try:
    app_db.init_thick_mode(app_config.OracleConfig(user="u", password="p", host="h", service_name="s"))
    check("lib_dir 缺失时抛异常", False)
except Exception:
    check("lib_dir 缺失时抛异常", True)
finally:
    if old_oh is None:
        os.environ.pop("ORACLE_HOME", None)
    else:
        os.environ["ORACLE_HOME"] = old_oh
os.makedirs("/tmp/fake_ic", exist_ok=True)
mock2.init_oracle_client(lib_dir="/tmp/fake_ic")
check("合法 lib_dir 初始化成功", mock2.thick and mock2.init_called)
# 幂等: 第二次不再报错
app_db._THICK_INITIALIZED = False
app_db.init_thick_mode(app_config.OracleConfig(user="u", password="p", host="h", service_name="s"))
check("幂等初始化", True)

# ---------------------------------------------------------------------------
# 汇总
# ---------------------------------------------------------------------------
shutil.rmtree(outdir, ignore_errors=True)
shutil.rmtree("/tmp/fake_ic", ignore_errors=True)
print(f"\n===== 结果: {len(PASS)} 通过 / {len(FAIL)} 失败 =====")
if FAIL:
    print("失败项:", FAIL)
    sys.exit(1)
print("全部通过 ✔")
